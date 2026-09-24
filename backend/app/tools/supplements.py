"""Supplement discovery independent of full-text XML, and bounded attachment parsing."""
from __future__ import annotations

import csv
import io
import re
import zipfile
from pathlib import PurePosixPath
from urllib.parse import urljoin, urlsplit, parse_qs, quote

import httpx
from bs4 import BeautifulSoup

from .http import ToolHttpError, USER_AGENT
from ..parsing.base import ParsedDocument, PdfParseError
from ..parsing.factory import get_pdf_parser

SUPPORTED = {'.pdf', '.xlsx', '.xls', '.csv', '.tsv', '.txt', '.docx'}
MAX_EXPANDED = 100 * 1024 * 1024
MAX_TEXT = 2_000_000


def attachment_links(html: str, base: str, source: str) -> list[dict]:
    soup = BeautifulSoup(html, 'html.parser')
    candidates = {}
    for a in soup.select('a[href]'):
        url = urljoin(base, a['href'])
        parts = urlsplit(url)
        label = a.get_text(' ', strip=True)
        container = a.find_parent(['section', 'div', 'table'])
        context = ' '.join(str(container.get(k, '')) for k in ('id', 'class')) if container else ''
        marker = f'{url} {label} {context}'.lower()
        if parts.scheme not in {'http', 'https'} or not any(x in marker for x in ('suppl', 'supporting', '/bin/', 'additional file')):
            continue
        filename = parse_qs(parts.query).get('file', [PurePosixPath(parts.path).name])[0]
        if PurePosixPath(filename.lower()).suffix not in SUPPORTED | {'.zip'}:
            continue
        candidates[url] = {'url': url, 'fileName': filename, 'description': label[:500], 'source': source}
    return list(candidates.values())


async def discovery_page(url: str) -> tuple[str, str]:
    async with httpx.AsyncClient(timeout=20, follow_redirects=True) as client:
        async with client.stream('GET', url, headers={'User-Agent': USER_AGENT}) as response:
            if response.status_code >= 400:
                raise ToolHttpError(f'HTTP {response.status_code} retrieving supplementary links.')
            body = bytearray()
            async for chunk in response.aiter_bytes():
                body.extend(chunk)
                if len(body) > 4 * 1024 * 1024:
                    raise ToolHttpError('Supplement discovery page exceeds 4 MB.')
            return body.decode('utf-8', errors='replace'), str(response.url)


async def discover(publication: dict) -> dict:
    candidates, checks = [], []
    pmcid, pmid, doi = (publication.get(key) for key in ('pmcid', 'pmid', 'doi'))
    pages = []
    if pmcid:
        pages.append(('pmc', f'https://pmc.ncbi.nlm.nih.gov/articles/{quote(pmcid, safe="")}/'))
    if doi:
        pages.append(('publisher', f'https://doi.org/{quote(doi, safe="/")}'))
    for source, url in pages:
        try:
            html, final = await discovery_page(url)
            found = attachment_links(html, final, source)
            candidates.extend(found)
            checks.append({'source': source, 'url': final, 'status': 'found' if found else 'not_found'})
        except (ToolHttpError, httpx.HTTPError) as error:
            checks.append({'source': source, 'url': url, 'status': 'discovery_failed', 'error': str(error)})
    # This service returns converted text, not original Excel formatting. Never
    # mistake its HTTP-200 error/help page for an actual supplementary document.
    if pmcid or pmid:
        url = f'https://www.ncbi.nlm.nih.gov/research/bionlp/RESTful/supplmat.cgi/bioc_xml/{quote(str(pmcid or pmid), safe="")}/list'
        try:
            html, _ = await discovery_page(url)
            soup = BeautifulSoup(html, 'xml')
            nums = [node.get_text(strip=True) for node in soup.find_all('num')]
            nums = list(dict.fromkeys(n for n in nums if n.isdigit()))
            for n in nums:
                candidates.append({'url': url.rsplit('/', 1)[0] + '/' + n,
                                   'fileName': f'PMC-supplement-{n}.xml', 'source': 'ncbi-bioc',
                                   'description': 'Converted supplementary text; original spreadsheet formatting is not preserved.'})
            checks.append({'source': 'ncbi-bioc', 'url': url,
                           'status': 'found' if nums else 'not_found',
                           'note': 'No indexed supplementary files returned.' if not nums else ''})
        except (ToolHttpError, httpx.HTTPError) as error:
            checks.append({'source': 'ncbi-bioc', 'url': url, 'status': 'discovery_failed', 'error': str(error)})
    candidates = list({item['url']: item for item in candidates}.values())
    return {'status': 'found' if candidates else ('discovery_failed' if any(c['status'] == 'discovery_failed' for c in checks) else 'not_found'),
            'candidates': candidates[:40], 'truncated': len(candidates) > 40, 'checks': checks}


def archive_members(data: bytes) -> list[dict]:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        infos = [item for item in archive.infolist() if not item.is_dir()]
        if len(infos) > 300 or sum(item.file_size for item in infos) > MAX_EXPANDED:
            raise PdfParseError('Archive exceeds 300 files or 100 MB expanded limit.')
        return [{'fileName': item.filename, 'bytes': item.file_size,
                 'supported': PurePosixPath(item.filename.lower()).suffix in SUPPORTED}
                for item in infos]


def table_text(rows) -> str:
    output, size = [], 0
    for i, row in enumerate(rows, 1):
        if i > 100_000:
            raise PdfParseError('Table exceeds 100,000 rows; provide a smaller sample-design table.')
        line = f'{i}\t' + '\t'.join(str(cell if cell is not None else '').replace('\n', '\\n').replace('\t', ' ') for cell in row)
        size += len(line)
        if size > MAX_TEXT:
            raise PdfParseError('Table exceeds 2 million characters; provide a smaller sample-design table.')
        output.append(line)
    return '\n'.join(output)


async def parse_attachment(data: bytes, filename: str, source: str = '') -> ParsedDocument:
    suffix = PurePosixPath(filename.lower()).suffix
    if data.startswith(b'%PDF-'):
        document = await get_pdf_parser().parse_bytes(data, filename)
        if not document.markdown.strip():
            raise PdfParseError('MinerU returned empty supplementary text.')
        return document
    if suffix == '.pdf':
        raise PdfParseError('Attachment is not a PDF (possibly an access-denied page).')
    sections = {}
    if source == 'ncbi-bioc':
        soup = BeautifulSoup(data, 'xml')
        text = '\n\n'.join(p.get_text(' ', strip=True) for p in soup.select('passage > text'))
        if not text.strip():
            raise PdfParseError('NCBI returned no supplementary passages.')
        sections = {'body': text}
    elif suffix == '.xlsx':
        import openpyxl
        archive_members(data)
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            for i, sheet in enumerate(workbook.worksheets, 1):
                sections[f'sheet {i}: {sheet.title.lower()}'] = table_text(sheet.iter_rows(values_only=True))
        finally:
            workbook.close()
    elif suffix == '.xls':
        import xlrd
        workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
        try:
            for i, sheet in enumerate(workbook.sheets(), 1):
                sections[f'sheet {i}: {sheet.name.lower()}'] = table_text(sheet.row_values(row) for row in range(sheet.nrows))
        finally:
            workbook.release_resources()
    elif suffix == '.docx':
        archive_members(data)
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            soup = BeautifulSoup(archive.read('word/document.xml'), 'xml')
            sections = {'body': '\n'.join(p.get_text() for p in soup.find_all('p'))}
    elif suffix in {'.csv', '.tsv', '.txt'}:
        text = data.decode('utf-8-sig')
        if re.search(r'<(?:!doctype\s+html|html|script)\b', text[:1000], re.I):
            raise PdfParseError('Attachment is an HTML page, not supplementary data.')
        sections = {'tables': table_text(csv.reader(io.StringIO(text), delimiter='\t' if suffix == '.tsv' else ','))} if suffix != '.txt' else {'body': text}
    else:
        raise PdfParseError(f'Unsupported supplementary format: {suffix or "unknown"}. Upload a PDF or tabular/text file.')
    markdown = '\n\n'.join(f'## {name}\n{text}' for name, text in sections.items() if text.strip())
    if not markdown.strip() or len(markdown) > MAX_TEXT:
        raise PdfParseError('Supplement is empty or exceeds 2 million characters.')
    return ParsedDocument(markdown, sections, parser='supplement-structured')


async def parse_uploaded_attachment(data: bytes, filename: str) -> ParsedDocument:
    """Combine a small uploaded archive, preserving member and section names."""
    if not filename.lower().endswith('.zip'):
        return await parse_attachment(data, filename)
    members = archive_members(data)
    supported = [m for m in members if m['supported']]
    if not supported or len(supported) > 20:
        raise PdfParseError('ZIP must contain 1–20 supported documents. Extract and upload the relevant table separately.')
    sections, total = {}, 0
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        for member in supported:
            name = member['fileName']
            try:
                doc = await parse_attachment(archive.read(name), name)
            except Exception as error:
                raise PdfParseError(f'Cannot parse ZIP member {name}: {error}. Extract and upload the relevant file separately.') from error
            for section, text in doc.sections.items():
                total += len(text)
                if total > MAX_TEXT:
                    raise PdfParseError('Combined archive exceeds 2 million characters. Upload the relevant table separately.')
                sections[f'{name} / {section}'] = text
    skipped = [m['fileName'] for m in members if not m['supported']]
    if skipped:
        sections['archive inventory (not evidence)'] = 'Unsupported files not parsed:\n' + '\n'.join(skipped)
    markdown = '\n\n'.join(f'## {name}\n{text}' for name, text in sections.items())
    return ParsedDocument(markdown, sections, parser='supplement-archive')


async def discover_pride(accession: str) -> dict:
    from .pride import normalize_accession, PRIDE_API_BASE
    from .http import get_json
    accession = normalize_accession(accession)
    payload = await get_json(f'{PRIDE_API_BASE}/projects/{accession}/files/all', timeout=60)
    entries = payload if isinstance(payload, list) else payload.get('_embedded', {}).get('files', payload.get('files', []))
    candidates = []
    for entry in entries or []:
        name = entry.get('fileName') or entry.get('name') or ''
        suffix = PurePosixPath(name.lower()).suffix
        # Raw acquisition archives are not sample-design attachments.
        if suffix not in SUPPORTED | {'.zip'} or (suffix == '.zip' and not re.search(r'suppl|support|table|design|metadata|sdrf', name, re.I)):
            continue
        for location in entry.get('publicFileLocations') or []:
            url = location.get('value', '') if isinstance(location, dict) else str(location)
            if url.startswith('ftp://ftp.pride.ebi.ac.uk/'):
                url = url.replace('ftp://ftp.pride.ebi.ac.uk/', 'https://ftp.pride.ebi.ac.uk/', 1)
            if urlsplit(url).scheme in {'http', 'https'}:
                candidates.append({'url': url, 'fileName': name, 'source': 'pride', 'accession': accession,
                                   'description': 'Project file candidate; verify contents and publication relationship before use.'})
                break
    return {'candidates': candidates[:40], 'truncated': len(candidates) > 40,
            'checks': [{'source': 'pride', 'accession': accession, 'status': 'found' if candidates else 'not_found'}]}
