import io
import zipfile
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
import openpyxl

from app.tools import supplements, registry
from app.session import SessionStore
from app.llm.setup_gate import SetupGate
from app.parsing.base import PdfParseError, ParsedDocument
from app.tools.http import ToolHttpError


@pytest.fixture
def store(monkeypatch):
    value = SessionStore()
    monkeypatch.setattr(registry, 'get_session_store', lambda: value)
    return value


def zip_bytes(files):
    output = io.BytesIO()
    with zipfile.ZipFile(output, 'w') as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    return output.getvalue()


def test_publisher_links_require_real_attachment():
    html = '''<a href="/action/downloadSupplement?doi=10.1/test&amp;file=SuppMat.zip">Supporting Table 1</a>
              <a href="/support">Support</a><a href="javascript:bad">Supplementary PDF</a>'''
    links = supplements.attachment_links(html, 'https://publisher.example/paper', 'publisher')
    assert len(links) == 1
    assert links[0]['fileName'] == 'SuppMat.zip'


async def test_discovery_keeps_failures_and_other_sources(monkeypatch):
    page = AsyncMock(side_effect=[
        ('<a href="/articles/PMC1/bin/table1.xlsx">Supplementary Table 1</a>', 'https://pmc.ncbi.nlm.nih.gov/articles/PMC1/'),
        ToolHttpError('HTTP 403'),
        ('<results><file><num>1</num></file></results>', 'https://ncbi.example/')])
    monkeypatch.setattr(supplements, 'discovery_page', page)
    result = await supplements.discover({'pmid': '1', 'pmcid': 'PMC1', 'doi': '10.123/test'})
    assert result['status'] == 'found'
    assert len(result['candidates']) == 2
    assert result['checks'][1]['status'] == 'discovery_failed'
    assert result['candidates'][1]['source'] == 'ncbi-bioc'


async def test_200_error_help_is_not_supplement(monkeypatch):
    monkeypatch.setattr(supplements, 'discovery_page', AsyncMock(return_value=('[Error] No result <html>Help</html>', 'https://ncbi.example/')))
    result = await supplements.discover({'pmid': '1'})
    assert result['status'] == 'not_found'
    assert not result['candidates']
    with pytest.raises(PdfParseError):
        await supplements.parse_attachment(b'<html>Access denied</html>', 'text.xml', 'ncbi-bioc')


async def test_abstract_read_cannot_unlock_gate(store, monkeypatch):
    publication = {'found': True, 'pmid': '123', 'doi': '10.123/test', 'abstract': 'Abstract only', 'url': 'https://doi.org/10.123/test'}
    monkeypatch.setattr(registry.literature, 'lookup_publication', AsyncMock(return_value=publication))
    result = await registry._find_publication({'pmid': '123'}, 's')
    doc_id = result['abstract']['documentId']
    assert 'sessionDocuments' not in result
    await registry._read_document({'documentId': doc_id}, 's')
    gate = SetupGate(store, 's', 'PXD1')
    gate.observe('get_pride_metadata', {'accession': 'PXD1'})
    gate.observe('find_publication', publication)
    assert gate.reason()


async def test_pubmed_fallback_without_xml(store, monkeypatch):
    monkeypatch.setattr(registry, 'get_text', AsyncMock(return_value='''<PubmedArticleSet><PubmedArticle><MedlineCitation><PMID>123</PMID><Article><ArticleTitle>Paper</ArticleTitle><Abstract><AbstractText Label="Methods">Two samples.</AbstractText></Abstract></Article></MedlineCitation><PubmedData><ArticleId IdType="doi">10.123/test</ArticleId></PubmedData></PubmedArticle></PubmedArticleSet>'''))
    result = await registry._get_abstract({'pmid': '123'}, 's')
    doc = store.get(result['documentId'])
    assert doc.metadata['evidenceKind'] == 'abstract'
    assert doc.document.sections['abstract'] == 'Methods: Two samples.'


async def test_zip_table_becomes_readable_matching_evidence(store, monkeypatch):
    workbook = openpyxl.Workbook()
    workbook.active.title = 'Sample Design'
    workbook.active.append(['Sample', 'Project'])
    workbook.active.append(['S1', 'PXD1'])
    output = io.BytesIO(); workbook.save(output)
    data = zip_bytes({'Table1.xlsx': output.getvalue(), 'image.png': b'not text'})
    url = 'https://publisher.example/SuppMat.zip'
    publication = {'found': True, 'doi': '10.123/test', 'pmid': '123', 'fullTextAvailable': False}
    monkeypatch.setattr(registry.literature, 'lookup_publication', AsyncMock(return_value=publication))
    monkeypatch.setattr(supplements, 'discover', AsyncMock(return_value={'status': 'found', 'candidates': [{'url': url, 'fileName': 'SuppMat.zip', 'source': 'publisher'}], 'checks': []}))
    monkeypatch.setattr(registry, 'cached_download', AsyncMock(return_value=(data, '/cache/zip')))
    await registry._find_supplements({'pmid': '123'}, 's')
    listing = await registry._get_supplement({'url': url}, 's')
    assert listing['status'] == 'downloaded'
    assert listing['members'][0]['fileName'] == 'Table1.xlsx'
    result = await registry._get_supplement({'url': url, 'member': 'Table1.xlsx'}, 's')
    assert result['evidenceKind'] == 'supplement'
    gate = SetupGate(store, 's', 'PXD1')
    gate.observe('get_pride_metadata', {'accession': 'PXD1'})
    gate.observe('find_publication', publication)
    assert gate.reason()
    read = await registry._read_document({'documentId': result['documentId']}, 's')
    assert read['sections']['sheet 1: sample design'] == '1\tSample\tProject\n2\tS1\tPXD1'
    assert gate.reason() is None
    assert (await registry._get_supplement({'url': url, 'member': 'Table1.xlsx'}, 's'))['cached']
    assert (await registry._get_supplement({'url': url}, 'other'))['status'] == 'not_discovered'
    monkeypatch.setattr(registry, 'cached_download', AsyncMock(side_effect=ToolHttpError('HTTP 403')))
    assert (await registry._get_supplement({'url': url}, 's'))['status'] == 'download_failed'


async def test_supplement_pdf_uses_mineru(monkeypatch):
    parse = AsyncMock(return_value=ParsedDocument('Full supplement', {'body': 'Full supplement'}, parser='mineru'))
    monkeypatch.setattr(supplements, 'get_pdf_parser', lambda: SimpleNamespace(parse_bytes=parse))
    result = await supplements.parse_attachment(b'%PDF-example', 'table.pdf')
    assert result.parser == 'mineru'
    parse.assert_awaited_once()


async def test_bad_attachment_and_archive_bounds():
    with pytest.raises(PdfParseError):
        await supplements.parse_attachment(b'<html>Denied</html>', 'table.csv')
    with pytest.raises(PdfParseError):
        supplements.archive_members(zip_bytes({f'{i}.txt': b'a' for i in range(301)}))
    document = await supplements.parse_attachment(b'Sample\tGroup\nS1\tControl', 'table.tsv')
    assert '2\tS1\tControl' in document.sections['tables']


async def test_abstract_cache_does_not_masquerade_as_full_text(store, monkeypatch):
    abstract = store.add_document('s', 'abstract.txt', ParsedDocument('Summary', {'abstract': 'Summary'}),
                                  metadata={'pmcid': 'PMC123', 'doi': '10.123/test', 'evidenceKind': 'abstract'})
    monkeypatch.setattr(registry.literature, 'fetch_full_text', AsyncMock(side_effect=ToolHttpError('No XML')))
    assert (await registry._get_full_text({'pmcid': 'PMC123'}, 's'))['status'] == 'download_failed'
    monkeypatch.setattr(registry, 'cached_download', AsyncMock(return_value=(b'%PDF-example', '/cache/paper')))
    parse = AsyncMock(return_value=ParsedDocument('Full paper', {'body': 'Full paper'}, parser='mineru'))
    monkeypatch.setattr(registry, 'get_pdf_parser', lambda: SimpleNamespace(parse_bytes=parse))
    result = await registry._parse_pdf_url({'url': 'https://publisher/paper.pdf', 'doi': '10.123/test'}, 's')
    assert result['documentId'] != abstract.document_id
    parse.assert_awaited_once()


@pytest.mark.asyncio
async def test_uploaded_archive_preserves_member_and_rows():
    doc = await supplements.parse_uploaded_attachment(zip_bytes({'Table1.tsv': 'sample\tfile\nS1\traw1', 'readme.bin': b'xx'}), 'support.zip')
    assert 'Table1.tsv / tables' in doc.sections
    assert '2\tS1\traw1' in doc.sections['Table1.tsv / tables']
    assert 'readme.bin' in doc.sections['archive inventory (not evidence)']


@pytest.mark.asyncio
async def test_direct_link_can_be_parsed_without_discovery(store, monkeypatch):
    monkeypatch.setattr(registry, 'cached_download', AsyncMock(return_value=(b'sample,file\nS1,raw1', '/tmp/test.raw')))
    result = await registry._get_supplement({'url': 'https://example.org/table.csv', 'userProvided': True}, 'direct')
    assert result['documentId']
    doc = store.get(result['documentId'])
    assert doc.metadata['source'] == 'user-link'
    assert not doc.metadata.get('doi')


@pytest.mark.asyncio
async def test_pride_attachment_discovery_ignores_raw_archives(monkeypatch):
    from app.tools import http
    monkeypatch.setattr(http, 'get_json', AsyncMock(return_value=[
        {'fileName': 'Exp3-sample1.zip', 'publicFileLocations': [{'value': 'ftp://ftp.pride.ebi.ac.uk/raw.zip'}]},
        {'fileName': 'Table1.xlsx', 'publicFileLocations': [{'value': 'ftp://ftp.pride.ebi.ac.uk/Table1.xlsx'}]},
    ]))
    result = await supplements.discover_pride('PXD001522')
    assert len(result['candidates']) == 1
    assert result['candidates'][0]['url'] == 'https://ftp.pride.ebi.ac.uk/Table1.xlsx'


@pytest.mark.asyncio
@pytest.mark.parametrize('filename,data', [('table.tsv', b'sample\tfile\nS1\traw1'), ('tables.zip', zip_bytes({'table.tsv': 'sample\tfile\nS1\traw1'}))])
async def test_uploaded_supplement_is_session_document(filename, data, monkeypatch):
    from app.routers import uploads
    from starlette.datastructures import UploadFile
    local = SessionStore()
    monkeypatch.setattr(uploads, 'get_session_store', lambda: local)
    result = await uploads.upload_pdf(sessionId='upload-test', file=UploadFile(io.BytesIO(data), filename=filename))
    doc = local.get(result.documentId)
    assert doc.metadata['evidenceKind'] == 'supplement'
    assert 'S1\traw1' in doc.document.markdown


@pytest.mark.asyncio
async def test_pride_only_discovery_does_not_assign_paper_identity(store, monkeypatch):
    monkeypatch.setattr(supplements, 'discover_pride', AsyncMock(return_value={
        'candidates': [{'url': 'https://example.org/table.tsv', 'fileName': 'table.tsv', 'source': 'pride', 'accession': 'PXD001522'}],
        'checks': [{'source': 'pride', 'status': 'found'}], 'truncated': False}))
    result = await registry._find_supplements({'accession': 'PXD001522'}, 'project')
    identifiers = store.pdf_identifiers('project', result['candidates'][0]['url'])
    assert identifiers['accession'] == 'PXD001522'
    assert 'doi' not in identifiers
